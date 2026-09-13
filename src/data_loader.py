"""
Auntie Aloha Dashboard - Data Loader & Pipeline
Parses Nabis remittances (2025-YTD), Nabis orders/inventory, QuickBooks exports, and budget models.
"""

import os
import glob
import re
import datetime
from typing import Dict, List, Optional, Tuple
import pandas as pd
import openpyxl


def parse_nabis_remittance_file(file_path: str) -> Tuple[pd.DataFrame, Dict]:
    """
    Parses a single Nabis remittance Excel workbook.
    Extracts raw line items and metadata (period dates, net remittance, fees).
    """
    fname = os.path.basename(file_path)
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    # 1. Detect target raw sheet
    raw_sheet_name = None
    for name in wb.sheetnames:
        lower = name.lower()
        if "raw" in lower or "auntie" in lower or "rrn" in lower:
            raw_sheet_name = name
            break
    if not raw_sheet_name:
        raw_sheet_name = wb.sheetnames[0]
        
    ws_raw = wb[raw_sheet_name]
    
    # Find header row
    hdr_row = 1
    headers = []
    for r in range(1, min(10, ws_raw.max_row + 1)):
        row_vals = [str(ws_raw.cell(r, c).value or "").strip() for c in range(1, ws_raw.max_column + 1)]
        if any("Paid Date" in v or "Order Number" in v for v in row_vals):
            hdr_row = r
            headers = row_vals
            break
            
    col_map = {h: i + 1 for i, h in enumerate(headers) if h}
    
    records = []
    for r in range(hdr_row + 1, ws_raw.max_row + 1):
        def get_val(col_name, default=None):
            idx = col_map.get(col_name)
            return ws_raw.cell(r, idx).value if idx else default

        def get_float(col_name):
            val = get_val(col_name, 0.0)
            try:
                if val is None:
                    return 0.0
                return float(val)
            except (ValueError, TypeError):
                return 0.0

        dispensary = str(get_val("Licensed Location Name", "") or "").strip()
        order_num = str(get_val("Order Number", "") or "").strip()
        inv_type = str(get_val("Invoice Type", "") or "").strip()
        paid_date = get_val("Paid Date")
        applied_date = get_val("Applied Date")
        
        nabis_collected = get_float("Nabis Collected")
        nabis_kept = get_float("Nabis Kept")
        nabis_remitted = get_float("Nabis Remitted to Brand")
        order_total = get_float("Order Total")
        inv_total = get_float("Invoice Total")
        
        # Categorize invoice / fee type
        inv_lower = inv_type.lower()
        if "due to processor" in inv_lower:
            fee_category = "Net Remittance to Brand"
        elif "fulfillment" in inv_lower:
            fee_category = "Distribution Fulfillment"
        elif "fuel" in inv_lower:
            fee_category = "Fuel Surcharge"
        elif "software" in inv_lower:
            fee_category = "Software Subscription"
        elif "case break" in inv_lower:
            fee_category = "Case Break Fee"
        elif "labeling" in inv_lower:
            fee_category = "Labeling Fee"
        elif "sample" in inv_lower or "trade" in inv_lower:
            fee_category = "Sample Order Fee"
        else:
            fee_category = "Other / Adjustments"

        # Format dates
        if isinstance(paid_date, datetime.datetime):
            paid_date = paid_date.date()
        if isinstance(applied_date, datetime.datetime):
            applied_date = applied_date.date()

        if order_num or inv_type or dispensary or nabis_collected != 0 or nabis_remitted != 0:
            records.append({
                "source_file": fname,
                "paid_date": paid_date,
                "applied_date": applied_date,
                "dispensary": dispensary,
                "order_number": order_num,
                "po_so_number": str(get_val("PO/SO Number", "") or "").strip(),
                "invoice_number": str(get_val("Invoice Number", "") or "").strip(),
                "invoice_type": inv_type,
                "fee_category": fee_category,
                "transaction_type": str(get_val("Transaction Type", "") or "").strip(),
                "order_total": order_total,
                "invoice_total": inv_total,
                "nabis_collected": nabis_collected,
                "nabis_kept": nabis_kept,
                "nabis_remitted": nabis_remitted,
            })
            
    df = pd.DataFrame(records)
    
    # 2. Extract metadata / period dates
    metadata = {
        "file": fname,
        "start_date": None,
        "end_date": None,
        "total_collected": float(df["nabis_collected"].sum()) if not df.empty else 0.0,
        "total_kept": float(df["nabis_kept"].sum()) if not df.empty else 0.0,
        "total_remitted": float(df["nabis_remitted"].sum()) if not df.empty else 0.0,
        "software_fee": float(df[df["fee_category"] == "Software Subscription"]["nabis_kept"].sum()) if not df.empty else 0.0,
    }
    
    # Check if 'Remittance Export' sheet exists for explicit dates
    if "Remittance Export" in wb.sheetnames:
        ws_meta = wb["Remittance Export"]
        for r in range(1, 10):
            val1 = str(ws_meta.cell(r, 2).value or "")
            if "Period Start Date" in val1:
                metadata["start_date"] = ws_meta.cell(r, 3).value
            elif "Period End Date" in val1:
                metadata["end_date"] = ws_meta.cell(r, 3).value
                
    # Fallback to filename regex for dates (e.g. 0106-01182025 or 2026-08-03_2026-08-15)
    if not metadata["start_date"]:
        match_iso = re.search(r"(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})", fname)
        if match_iso:
            metadata["start_date"] = match_iso.group(1)
            metadata["end_date"] = match_iso.group(2)
        else:
            match_mdy = re.search(r"(\d{2})(\d{2})-(\d{2})(\d{2})(\d{4})", fname)
            if match_mdy:
                m1, d1, m2, d2, y = match_mdy.groups()
                metadata["start_date"] = f"{y}-{m1}-{d1}"
                metadata["end_date"] = f"{y}-{m2}-{d2}"
                
    if metadata["start_date"] and isinstance(metadata["start_date"], (datetime.datetime, datetime.date)):
        metadata["start_date"] = str(metadata["start_date"])[:10]
    if metadata["end_date"] and isinstance(metadata["end_date"], (datetime.datetime, datetime.date)):
        metadata["end_date"] = str(metadata["end_date"])[:10]

    return df, metadata


def load_all_nabis_remittances(base_dir: str) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """
    Scans base_dir for all Nabis remittance XLSX files (including 2025, 2026 subdirs).
    Returns (raw_transactions_df, biweekly_summary_df).
    """
    pattern = os.path.join(base_dir, "**", "*.xlsx")
    files = glob.glob(pattern, recursive=True)
    # Filter out cashflow model itself
    remittance_files = [f for f in files if "cashflow" not in os.path.basename(f).lower() and not os.path.basename(f).startswith("~$")]
    
    all_rows = []
    summaries = []
    
    for f in sorted(remittance_files):
        try:
            df, meta = parse_nabis_remittance_file(f)
            if not df.empty:
                all_rows.append(df)
            summaries.append(meta)
        except Exception as e:
            print(f"Warning: Failed to parse {f}: {e}")
            
    if all_rows:
        raw_df = pd.concat(all_rows, ignore_index=True)
    else:
        raw_df = pd.DataFrame()
        
    summary_df = pd.DataFrame(summaries)
    if not summary_df.empty and "start_date" in summary_df.columns:
        summary_df["start_date"] = pd.to_datetime(summary_df["start_date"], errors="coerce")
        summary_df = summary_df.sort_values("start_date").reset_index(drop=True)
        summary_df["realization_pct"] = (summary_df["total_remitted"] / summary_df["total_collected"].replace(0, float("nan"))) * 100
        summary_df["realization_pct"] = summary_df["realization_pct"].fillna(0)
        
    return raw_df, summary_df


def load_budget_cashflow_model(file_path: str) -> Dict[str, pd.DataFrame]:
    """
    Loads baseline budget, income, expenses, and summary projections from
    'Auntie Aloha Cashflow Model.xlsx'.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    month_names = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    
    # 1. Income tab
    income_records = []
    if "Income" in wb.sheetnames:
        ws_inc = wb["Income"]
        for r in range(2, ws_inc.max_row + 1):
            category = str(ws_inc.cell(r, 2).value or "").strip()
            item_name = str(ws_inc.cell(r, 5).value or "").strip()
            if not item_name and not category:
                continue
            name = item_name if item_name else category
            values = [float(ws_inc.cell(r, c).value or 0.0) for c in range(6, 18)]
            if any(v != 0 for v in values) or item_name:
                row_dict = {"Category": category, "Stream": name}
                for m_idx, m_name in enumerate(month_names):
                    row_dict[m_name] = values[m_idx] if m_idx < len(values) else 0.0
                income_records.append(row_dict)
                
    income_df = pd.DataFrame(income_records)
    
    # 2. Expenses tab
    expense_records = []
    if "Expenses" in wb.sheetnames:
        ws_exp = wb["Expenses"]
        for r in range(2, ws_exp.max_row + 1):
            category = str(ws_exp.cell(r, 2).value or "").strip()
            item_name = str(ws_exp.cell(r, 5).value or "").strip()
            if not item_name and not category:
                continue
            name = item_name if item_name else category
            values = [float(ws_exp.cell(r, c).value or 0.0) for c in range(6, 18)]
            if any(v != 0 for v in values) or item_name:
                row_dict = {"Category": category, "Expense": name}
                for m_idx, m_name in enumerate(month_names):
                    row_dict[m_name] = values[m_idx] if m_idx < len(values) else 0.0
                expense_records.append(row_dict)
                
    expense_df = pd.DataFrame(expense_records)
    
    # 3. Starting Balance
    starting_balance = 14000.0
    if "Setup" in wb.sheetnames:
        ws_setup = wb["Setup"]
        for r in range(1, 15):
            val = str(ws_setup.cell(r, 2).value or "")
            if "Starting balance" in val:
                try:
                    starting_balance = float(ws_setup.cell(r, 3).value or 14000.0)
                except:
                    pass
                    
    return {
        "starting_balance": starting_balance,
        "income": income_df,
        "expenses": expense_df,
    }


def get_default_inventory_data() -> pd.DataFrame:
    """
    Returns baseline/current inventory tracking dataframe for Auntie Aloha SKUs
    (Gummies Distillate, Rosin, and flavors) until live Nabis inventory CSV is uploaded.
    """
    records = [
        {
            "sku": "AA-GUM-DIST-100MG",
            "product_name": "Auntie Aloha Distillate Gummies 100mg",
            "category": "Gummies - Distillate",
            "warehouse": "Nabis Oakland",
            "units_on_hand": 1840,
            "units_reserved": 320,
            "units_available": 1520,
            "weekly_velocity": 245,
            "reorder_threshold": 500,
            "batch_cost": 2.20,
            "wholesale_price": 7.50,
            "manufacturer": "Smoakland",
        },
        {
            "sku": "AA-GUM-DIST-SOU-100MG",
            "product_name": "Auntie Aloha Sour Tropical Distillate 100mg",
            "category": "Gummies - Distillate",
            "warehouse": "Nabis Los Angeles",
            "units_on_hand": 960,
            "units_reserved": 180,
            "units_available": 780,
            "weekly_velocity": 130,
            "reorder_threshold": 300,
            "batch_cost": 2.20,
            "wholesale_price": 7.50,
            "manufacturer": "Smoakland",
        },
        {
            "sku": "AA-GUM-ROSN-100MG",
            "product_name": "Auntie Aloha Live Rosin Gummies 100mg",
            "category": "Gummies - Solventless Rosin",
            "warehouse": "Nabis Oakland",
            "units_on_hand": 540,
            "units_reserved": 90,
            "units_available": 450,
            "weekly_velocity": 95,
            "reorder_threshold": 250,
            "batch_cost": 3.80,
            "wholesale_price": 10.00,
            "manufacturer": "MyGreen Network",
        },
        {
            "sku": "AA-GUM-ROSN-LILIKOI",
            "product_name": "Auntie Aloha Lilikoi Passionfruit Rosin",
            "category": "Gummies - Solventless Rosin",
            "warehouse": "Nabis Los Angeles",
            "units_on_hand": 310,
            "units_reserved": 40,
            "units_available": 270,
            "weekly_velocity": 65,
            "reorder_threshold": 180,
            "batch_cost": 3.80,
            "wholesale_price": 10.00,
            "manufacturer": "MyGreen Network",
        },
    ]
    df = pd.DataFrame(records)
    df["days_of_supply"] = (df["units_available"] / (df["weekly_velocity"] / 7.0)).round(1)
    df["weeks_of_supply"] = (df["units_available"] / df["weekly_velocity"]).round(1)
    df["status"] = df["days_of_supply"].apply(
        lambda d: "🔴 Critical Reorder" if d < 21 else ("🟡 Reorder Soon" if d < 35 else "🟢 Healthy")
    )
    return df


def parse_nabis_inventory_export(file_or_buffer) -> pd.DataFrame:
    """
    Parses a Nabis inventory CSV or Excel export.
    Uses smart fuzzy column matching to handle various Nabis report versions:
    - Product / Item Name
    - SKU
    - Units Available / Total On Hand / Reserved
    - Warehouse Location (Oakland, Los Angeles)
    """
    try:
        if hasattr(file_or_buffer, "name"):
            fname = file_or_buffer.name.lower()
            if fname.endswith(".csv"):
                df_raw = pd.read_csv(file_or_buffer)
            else:
                df_raw = pd.read_excel(file_or_buffer)
        elif isinstance(file_or_buffer, str):
            if file_or_buffer.lower().endswith(".csv"):
                df_raw = pd.read_csv(file_or_buffer)
            else:
                df_raw = pd.read_excel(file_or_buffer)
        else:
            df_raw = pd.DataFrame(file_or_buffer)

        # Build clean column map
        col_map = {orig: str(orig).strip().lower() for orig in df_raw.columns}

        def find_col(patterns):
            for orig, clean in col_map.items():
                if any(p in clean for p in patterns):
                    return orig
            return None

        col_product = find_col(["product name", "item name", "product", "item", "description", "title"])
        col_sku = find_col(["sku", "item code", "code", "upc"])
        col_avail = find_col(["available", "sellable", "salable", "units available"])
        col_total = find_col(["total on hand", "on hand", "units on hand", "total units", "quantity", "qty"])
        col_reserved = find_col(["reserved", "allocated", "hold"])
        col_whs = find_col(["warehouse", "facility", "location", "hub", "site"])

        records = []
        for idx, row in df_raw.iterrows():
            prod = str(row[col_product]).strip() if col_product else f"SKU {idx+1}"
            if not prod or prod.lower() == "nan" or prod.lower() == "total":
                continue

            sku = str(row[col_sku]).strip() if col_sku else f"AA-SKU-{idx+1}"
            whs = str(row[col_whs]).strip() if col_whs else "Nabis Warehouse"

            def get_num(col):
                if col:
                    try:
                        v = float(row[col])
                        return v if not pd.isna(v) else 0.0
                    except:
                        return 0.0
                return 0.0

            total_units = get_num(col_total)
            avail_units = get_num(col_avail)
            res_units = get_num(col_reserved)
            if avail_units == 0 and total_units > 0:
                avail_units = max(0.0, total_units - res_units)
            if total_units == 0 and avail_units > 0:
                total_units = avail_units + res_units

            # Determine category & manufacturer
            p_lower = prod.lower()
            if "rosin" in p_lower:
                cat = "Gummies - Solventless Rosin"
                mfg = "MyGreen Network"
                cost = 3.80
                whs_price = 10.00
                default_vel = 80.0
            else:
                cat = "Gummies - Distillate"
                mfg = "Smoakland"
                cost = 2.20
                whs_price = 7.50
                default_vel = 180.0

            records.append({
                "sku": sku,
                "product_name": prod,
                "category": cat,
                "warehouse": whs,
                "units_on_hand": total_units,
                "units_reserved": res_units,
                "units_available": avail_units,
                "weekly_velocity": default_vel,
                "reorder_threshold": round(default_vel * 2.0),
                "batch_cost": cost,
                "wholesale_price": whs_price,
                "manufacturer": mfg,
            })

        return pd.DataFrame(records)
    except Exception as e:
        print(f"Error parsing Nabis inventory export: {e}")
        return pd.DataFrame()

